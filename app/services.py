import hashlib
import json
import re
import shutil
import sqlite3
import uuid
from datetime import datetime, timezone
from pathlib import Path

from .db import get_db, utc_now


ALLOWED_SEMANTICS = {"unknown", "asset", "serial"}
CODE_RE = re.compile(r"^[^\s|]+$")
STATUS_LABELS = {
    "created": "未完成",
    "started": "未完成",
    "finished": "已完成",
    "cancelled": "已取消",
}
MODEL_SERIAL_PREFIXES = {"HR255": "E", "RC255": "ER", "AR210": "C", "AR170": "B", "RC210": "CR"}


def infer_serial_prefix(model_name, model_code):
    for value in (model_code, model_name):
        key = "".join(char for char in (value or "").upper() if char.isalnum())
        if key in MODEL_SERIAL_PREFIXES:
            return MODEL_SERIAL_PREFIXES[key]
    return ""


def new_id():
    return str(uuid.uuid4())


def status_label(status):
    return STATUS_LABELS.get(status, status or "-")


def format_duration(started_at, finished_at, status=None):
    if not started_at:
        return "-"
    if not finished_at:
        return "進行中" if status == "started" else "-"
    try:
        started = datetime.fromisoformat(started_at)
        finished = datetime.fromisoformat(finished_at)
    except (TypeError, ValueError):
        return "-"
    seconds = max(0, int((finished - started).total_seconds()))
    days, remainder = divmod(seconds, 86400)
    hours, remainder = divmod(remainder, 3600)
    minutes, seconds = divmod(remainder, 60)
    duration = f"{hours:02d}:{minutes:02d}:{seconds:02d}"
    return f"{days}日 {duration}" if days else duration


def validate_text(value, field, required=True):
    value = (value or "").strip()
    if required and not value:
        raise ValueError(f"{field}不可為空")
    if "|" in value:
        raise ValueError(f"{field}不可包含 |")
    return value


def validate_code(value, field):
    value = validate_text(value, field)
    if not CODE_RE.match(value):
        raise ValueError(f"{field}格式不正確")
    return value


def add_audit(action, entity_type, entity_id=None, details=None, operator_name=None):
    db = get_db()
    db.execute(
        "INSERT INTO audit_log(log_id, action, entity_type, entity_id, details, operator_name, created_at) VALUES (?,?,?,?,?,?,?)",
        (new_id(), action, entity_type, entity_id, json.dumps(details or {}, ensure_ascii=False), operator_name, utc_now()),
    )


def get_setting(key, default=None):
    row = get_db().execute("SELECT value FROM app_settings WHERE key=?", (key,)).fetchone()
    return row["value"] if row else default


def set_setting(key, value, operator_name="system"):
    if key == "number_semantics" and value not in ALLOWED_SEMANTICS:
        raise ValueError("number_semantics 必須是 unknown、asset 或 serial")
    db = get_db()
    db.execute(
        "INSERT INTO app_settings(key,value,updated_at) VALUES(?,?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_at=excluded.updated_at",
        (key, value, utc_now()),
    )
    add_audit("update_setting", "app_setting", key, {"value": value}, operator_name)
    db.commit()


def list_models(active_only=False):
    sql = "SELECT * FROM model_master"
    if active_only:
        sql += " WHERE is_active=1"
    return get_db().execute(sql + " ORDER BY sort_order, model_name").fetchall()


def list_assembly_workers(active_only=False):
    sql = "SELECT * FROM assembly_worker"
    if active_only:
        sql += " WHERE is_active=1"
    return get_db().execute(sql + " ORDER BY name COLLATE NOCASE").fetchall()


def get_assembly_worker(name):
    name = validate_text(name, "組裝人員")
    return get_db().execute(
        "SELECT * FROM assembly_worker WHERE name=? COLLATE NOCASE AND is_active=1",
        (name,),
    ).fetchone()


def create_assembly_worker(name, changed_by="admin"):
    name = validate_text(name, "組裝人員")
    db = get_db()
    now = utc_now()
    try:
        cur = db.execute(
            "INSERT INTO assembly_worker(name,created_at,updated_at) VALUES(?,?,?)",
            (name, now, now),
        )
    except sqlite3.IntegrityError as exc:
        raise ValueError("此組裝人員已存在") from exc
    add_audit("create", "assembly_worker", str(cur.lastrowid), {"name": name}, changed_by)
    db.commit()
    return cur.lastrowid


def set_assembly_worker_active(worker_id, is_active, changed_by="admin"):
    db = get_db()
    worker = db.execute("SELECT name FROM assembly_worker WHERE id=?", (worker_id,)).fetchone()
    if not worker:
        raise ValueError("找不到組裝人員")
    db.execute(
        "UPDATE assembly_worker SET is_active=?,updated_at=? WHERE id=?",
        (int(bool(is_active)), utc_now(), worker_id),
    )
    add_audit("update", "assembly_worker", str(worker_id), {"name": worker["name"], "is_active": bool(is_active)}, changed_by)
    db.commit()


def create_model(model_name, model_code, description="", operator_name="system", serial_prefix=None):
    model_name = validate_text(model_name, "機型名稱")
    model_code = validate_code(model_code, "機型代碼")
    serial_prefix = validate_code(serial_prefix, "製號字首") if serial_prefix else infer_serial_prefix(model_name, model_code)
    db = get_db()
    now = utc_now()
    try:
        cur = db.execute(
            "INSERT INTO model_master(model_name,model_code,serial_prefix,description,created_at,updated_at) VALUES(?,?,?,?,?,?)",
            (model_name, model_code, serial_prefix, description.strip(), now, now),
        )
    except sqlite3.IntegrityError as exc:
        raise ValueError("機型代碼已存在") from exc
    add_audit("create", "model_master", str(cur.lastrowid), {"model_code": model_code}, operator_name)
    db.commit()
    return cur.lastrowid


def update_model(model_id, model_name, model_code, description, is_active, operator_name="system", serial_prefix=None):
    model_name = validate_text(model_name, "機型名稱")
    model_code = validate_code(model_code, "機型代碼")
    serial_prefix = validate_code(serial_prefix, "製號字首") if serial_prefix else infer_serial_prefix(model_name, model_code)
    db = get_db()
    try:
        db.execute("UPDATE model_master SET model_name=?,model_code=?,serial_prefix=?,description=?,is_active=?,updated_at=? WHERE id=?", (model_name, model_code, serial_prefix, description.strip(), int(bool(is_active)), utc_now(), model_id))
    except sqlite3.IntegrityError as exc:
        raise ValueError("機型代碼已存在") from exc
    add_audit("update", "model_master", str(model_id), {"model_code": model_code, "is_active": bool(is_active)}, operator_name)
    db.commit()


def list_machines(model_id=None, active_only=False):
    clauses, params = [], []
    if model_id:
        clauses.append("model_id=?")
        params.append(model_id)
    if active_only:
        clauses.append("is_active=1")
    sql = "SELECT * FROM machine_master"
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    return get_db().execute(sql + " ORDER BY sort_order, number_value, display_no", params).fetchall()


def machine_number_prefix(model_id):
    machines = get_db().execute(
        "SELECT display_no FROM machine_master WHERE model_id=? AND is_active=1 ORDER BY id",
        (model_id,),
    ).fetchall()
    prefixes = set()
    for machine in machines:
        prefix, separator, suffix = machine["display_no"].rpartition("-")
        if not separator or not prefix or not suffix.isdigit():
            return ""
        prefixes.add(prefix + separator)
    return next(iter(prefixes)) if len(prefixes) == 1 else ""


def machine_ids_from_numbers(model_id, machine_numbers):
    tokens = [token.strip() for token in (machine_numbers or "").split(",")]
    if not tokens or any(not token for token in tokens):
        raise ValueError("請以逗號分隔輸入機台編號")
    normalized = [token.casefold() for token in tokens]
    if len(normalized) != len(set(normalized)):
        raise ValueError("機台編號不可重複")

    available = get_db().execute(
        "SELECT id, display_no FROM machine_master WHERE model_id=? AND is_active=1",
        (model_id,),
    ).fetchall()
    machines_by_number = {row["display_no"].casefold(): row["id"] for row in available}
    prefix = machine_number_prefix(model_id)
    machine_numbers = [prefix + token if prefix and token.isdigit() else token for token in tokens]
    invalid = [number for number, token in zip(machine_numbers, tokens) if number.casefold() not in machines_by_number or (prefix and not token.isdigit())]
    if invalid:
        raise ValueError("機台不存在、未啟用或不屬於所選機型：" + ", ".join(invalid))
    return [machines_by_number[number.casefold()] for number in machine_numbers]


def _machine_semantics(value):
    value = value or get_setting("number_semantics", "unknown")
    if value not in ALLOWED_SEMANTICS:
        raise ValueError("編號語意必須是 unknown、asset 或 serial")
    return value


def create_machine(model_id, display_no, number_prefix="", number_value=None, number_width=None, semantics=None, operator_name="system"):
    display_no = validate_code(display_no, "機台編號")
    semantics = _machine_semantics(semantics)
    db = get_db()
    now = utc_now()
    try:
        cur = db.execute(
            "INSERT INTO machine_master(model_id,display_no,number_prefix,number_value,number_width,number_semantics,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)",
            (model_id, display_no, number_prefix.strip(), number_value, number_width, semantics, now, now),
        )
    except sqlite3.IntegrityError as exc:
        raise ValueError("同一機型的機台編號不可重複") from exc
    add_audit("create", "machine_master", str(cur.lastrowid), {"display_no": display_no, "number_semantics": semantics}, operator_name)
    db.commit()
    return cur.lastrowid


def create_machine_range(model_id, prefix, start, end, width, semantics="unknown", operator_name="system"):
    prefix = validate_code(prefix, "機台前綴")
    semantics = _machine_semantics(semantics)
    if start < 0 or end < start or width < 1:
        raise ValueError("機台範圍或固定寬度不正確")
    db = get_db()
    now = utc_now()
    ids = []
    try:
        for value in range(start, end + 1):
            display_no = f"{prefix}-{value:0{width}d}"
            cur = db.execute("INSERT INTO machine_master(model_id,display_no,number_prefix,number_value,number_width,number_semantics,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)", (model_id, display_no, prefix, value, width, semantics, now, now))
            ids.append(cur.lastrowid)
            add_audit("create", "machine_master", str(cur.lastrowid), {"display_no": display_no, "number_semantics": semantics}, operator_name)
        db.commit()
    except sqlite3.IntegrityError as exc:
        db.rollback()
        raise ValueError("同一機型的機台編號不可重複") from exc
    return ids


def update_machine(machine_id, display_no, semantics, is_active, operator_name="system"):
    display_no = validate_code(display_no, "機台編號")
    semantics = _machine_semantics(semantics)
    db = get_db()
    try:
        db.execute("UPDATE machine_master SET display_no=?,number_semantics=?,is_active=?,updated_at=? WHERE id=?", (display_no, semantics, int(bool(is_active)), utc_now(), machine_id))
    except sqlite3.IntegrityError as exc:
        raise ValueError("同一機型的機台編號不可重複") from exc
    add_audit("update", "machine_master", str(machine_id), {"display_no": display_no, "number_semantics": semantics, "is_active": bool(is_active)}, operator_name)
    db.commit()


def get_batch(batch_id):
    return get_db().execute("SELECT * FROM production_batch WHERE batch_id=?", (batch_id,)).fetchone()


def get_job(job_run_id):
    return get_db().execute("SELECT * FROM job_run WHERE job_run_id=?", (job_run_id,)).fetchone()


def create_batch(model_id, manufacturing_year, starting_number, product_count, operator_name):
    try:
        manufacturing_year = int(manufacturing_year)
        starting_number = int(starting_number)
        product_count = int(product_count)
    except (TypeError, ValueError) as exc:
        raise ValueError("製作年份、起始流水號與台數都必須是正整數") from exc
    if not 2000 <= manufacturing_year <= 2099 or starting_number < 1 or product_count < 1:
        raise ValueError("製作年份需為 2000–2099，起始流水號與台數必須是正整數")
    operator_name = validate_text(operator_name, "組裝人員")
    db = get_db()
    model = db.execute("SELECT * FROM model_master WHERE id=? AND is_active=1", (model_id,)).fetchone()
    if not model:
        raise ValueError("機型不存在或已停用")
    serial_prefix = model["serial_prefix"] or infer_serial_prefix(model["model_name"], model["model_code"])
    if not serial_prefix:
        raise ValueError("此機型尚未設定製號字首，請至機型主檔設定")
    year_marker = f"{(manufacturing_year + 1) % 100:02d}"
    serials = [f"{serial_prefix}{year_marker}-{number:02d}" for number in range(starting_number, starting_number + product_count)]
    placeholders = ",".join("?" for _ in serials)
    used = db.execute(f"SELECT product_serial FROM job_run WHERE product_serial IN ({placeholders})", serials).fetchall()
    legacy_used = db.execute(f"SELECT machine_no_snapshot FROM job_run WHERE machine_no_snapshot IN ({placeholders})", serials).fetchall()
    if used or legacy_used:
        duplicates = {row[0] for row in [*used, *legacy_used]}
        raise ValueError("以下製號已存在，請更換起始流水號：" + ", ".join(sorted(duplicates)))
    now = utc_now()
    batch_id = new_id()
    batch_no = "B-" + datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S") + "-" + batch_id[:6].upper()
    db.execute("INSERT INTO production_batch(batch_id,batch_no,selection_mode,operator_name,created_at) VALUES(?,?,?,?,?)", (batch_id, batch_no, "serial_range", operator_name, now))
    job_ids = []
    for serial in serials:
        job_id = new_id()
        job_ids.append(job_id)
        db.execute("INSERT INTO job_run(job_run_id,batch_id,model_id,machine_id,product_serial,manufacturing_year,model_name_snapshot,model_code_snapshot,machine_no_snapshot,number_semantics_snapshot,planned_qty,status,operator_name,created_at,updated_at) VALUES(?,?,?,NULL,?,?,?,?,?,'serial',NULL,'created',?,?,?)", (job_id, batch_id, model_id, serial, manufacturing_year, model["model_name"], model["model_code"], serial, operator_name, now, now))
    add_audit("create", "production_batch", batch_id, {"job_count": len(job_ids), "manufacturing_year": manufacturing_year, "serials": serials}, operator_name)
    db.commit()
    return batch_id, job_ids


def checksum_for(parts):
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()[:8].upper()


def compact_batch_token(batch_id):
    return batch_id.replace("-", "").upper()[:8]


def make_payload(kind, job):
    parts = ["INTJR", "V1", kind, job["batch_id"], job["job_run_id"], job["model_code_snapshot"], job["machine_no_snapshot"]]
    return "|".join(parts + [checksum_for(parts)])


def make_batch_payload(kind, batch, jobs):
    prefix = "I1" + ("S" if kind == "START" else "F")
    token = compact_batch_token(batch["batch_id"])
    return prefix + token + checksum_for([prefix, token])[:4]


def parse_payload(payload):
    raw = (payload or "").strip()
    if len(raw) == 15 and raw[:2] == "I1" and raw[2] in {"S", "F"}:
        prefix, token, check = raw[:3], raw[3:11], raw[11:]
        if checksum_for([prefix, token])[:4] != check:
            raise ValueError("條碼 checksum 不正確")
        return {"system": "INTJR", "version": "1", "type": "START" if raw[2] == "S" else "FINISH", "batch_token": token, "checksum": check, "scope": "batch"}
    parts = raw.split("|")
    if len(parts) == 5 and parts[0] == "INTJR" and parts[1] in {"1", "V1"}:
        if parts[2] not in {"S", "F", "START", "FINISH"} or checksum_for(parts[:-1]) != parts[-1]:
            raise ValueError("條碼 checksum 或類型不正確")
        return {"system": parts[0], "version": parts[1], "type": {"S": "START", "F": "FINISH"}.get(parts[2], parts[2]), "batch_token": parts[3], "checksum": parts[4], "scope": "batch"}
    if len(parts) != 8 or parts[0] != "INTJR":
        raise ValueError("無效條碼格式")
    if parts[1] != "V1":
        raise ValueError("不支援的條碼版本")
    if parts[2] not in {"START", "FINISH"}:
        raise ValueError("不支援的條碼類型")
    if checksum_for(parts[:-1]) != parts[-1]:
        raise ValueError("條碼 checksum 不正確")
    return {"system": parts[0], "version": parts[1], "type": parts[2], "batch_id": parts[3], "job_run_id": parts[4], "model_code": parts[5], "machine_no": parts[6], "checksum": parts[7], "scope": "job_run"}


def ensure_barcodes(batch_id):
    db = get_db()
    jobs = db.execute("SELECT * FROM job_run WHERE batch_id=? ORDER BY id", (batch_id,)).fetchall()
    now = utc_now()
    batch = get_batch(batch_id)
    if batch["selection_mode"] == "serial_range" or get_setting("batch_scan_mode", "batch") == "batch":
        for kind in ("START", "FINISH"):
            payload = make_batch_payload(kind, batch, jobs)
            db.execute("INSERT OR IGNORE INTO barcode_record(batch_id,barcode_scope,barcode_type,payload,created_at) VALUES(?,?,?,?,?)", (batch_id, "batch", kind, payload, now))
    else:
        for job in jobs:
            for kind in ("START", "FINISH"):
                payload = make_payload(kind, job)
                db.execute("INSERT OR IGNORE INTO barcode_record(job_run_id,barcode_scope,barcode_type,payload,created_at) VALUES(?,?,?,?,?)", (job["job_run_id"], "job_run", kind, payload, now))
    db.commit()
    return jobs


def scan(payload, operator_name="scanner"):
    db = get_db()
    event_id = new_id()
    scanned_at = utc_now()
    try:
        parsed = parse_payload(payload)
        kind = parsed["type"]
        if parsed.get("scope") == "batch":
            candidates = db.execute("SELECT * FROM production_batch WHERE upper(replace(batch_id, '-', '')) LIKE ?", (parsed["batch_token"] + "%",)).fetchall()
            if len(candidates) != 1:
                raise ValueError("找不到相符的報工批次")
            batch_id = candidates[0]["batch_id"]
            jobs = db.execute("SELECT * FROM job_run WHERE batch_id=? ORDER BY id", (batch_id,)).fetchall()
            if not jobs:
                raise ValueError("找不到相符的報工批次")
            if kind == "START":
                pending = [job for job in jobs if job["status"] == "created"]
                if pending:
                    db.executemany("UPDATE job_run SET status='started',started_at=?,updated_at=? WHERE job_run_id=?", [(scanned_at, scanned_at, job["job_run_id"]) for job in pending])
                    result, message = "success", f"整批 START 成功：{len(pending)} 台產品"
                elif all(job["status"] == "started" for job in jobs):
                    result, message = "duplicate", "此批次已完成 START"
                else:
                    result, message = "invalid", "此批次已完成或已取消"
            else:
                started = [job for job in jobs if job["status"] == "started"]
                created = [job for job in jobs if job["status"] == "created"]
                if created:
                    result, message = "invalid", "批次內仍有產品未 START，不能 FINISH"
                elif started:
                    db.executemany("UPDATE job_run SET status='finished',actual_qty=CASE WHEN product_serial IS NULL THEN planned_qty ELSE actual_qty END,finished_at=?,updated_at=? WHERE job_run_id=?", [(scanned_at, scanned_at, job["job_run_id"]) for job in started])
                    result, message = "success", f"整批 FINISH 成功：{len(started)} 台產品"
                elif all(job["status"] == "finished" for job in jobs):
                    result, message = "duplicate", "此批次已完成 FINISH"
                else:
                    result, message = "invalid", "此批次已取消或沒有可完成的產品"
            db.execute("INSERT INTO scan_event(event_id,payload,barcode_type,job_run_id,result,message,scanned_at,operator_name) VALUES(?,?,?,?,?,?,?,?)", (event_id, payload, kind, None, result, message, scanned_at, operator_name))
            db.commit()
            return {"event_id": event_id, "result": result, "message": message, "scanned_at": scanned_at}

        job = get_job(parsed["job_run_id"])
        if not job or job["batch_id"] != parsed["batch_id"] or job["model_code_snapshot"] != parsed["model_code"] or job["machine_no_snapshot"] != parsed["machine_no"]:
            raise ValueError("找不到相符的報工資料")
        if kind == "START":
            if job["status"] == "created":
                db.execute("UPDATE job_run SET status='started',started_at=?,updated_at=? WHERE job_run_id=?", (scanned_at, scanned_at, job["job_run_id"]))
                result, message = "success", f"START 成功：{job['machine_no_snapshot']}"
            elif job["status"] == "started":
                result, message = "duplicate", f"已於 {job['started_at']} START"
            else:
                result, message = "invalid", "此報工已完成或已取消"
        else:
            if job["status"] == "started":
                db.execute("UPDATE job_run SET status='finished',actual_qty=CASE WHEN product_serial IS NULL THEN planned_qty ELSE actual_qty END,finished_at=?,updated_at=? WHERE job_run_id=?", (scanned_at, scanned_at, job["job_run_id"]))
                result, message = "success", f"FINISH 成功：{job['machine_no_snapshot']}"
            elif job["status"] == "created":
                result, message = "invalid", "必須先掃描 START"
            elif job["status"] == "finished":
                result, message = "duplicate", f"已於 {job['finished_at']} 完成"
            else:
                result, message = "invalid", "此報工已取消"
        db.execute("INSERT INTO scan_event(event_id,payload,barcode_type,job_run_id,result,message,scanned_at,operator_name) VALUES(?,?,?,?,?,?,?,?)", (event_id, payload, kind, job["job_run_id"], result, message, scanned_at, operator_name))
    except ValueError as exc:
        result, message = "invalid", str(exc)
        db.execute("INSERT INTO scan_event(event_id,payload,result,message,scanned_at,operator_name) VALUES(?,?,?,?,?,?)", (event_id, payload, result, message, scanned_at, operator_name))
    db.commit()
    return {"event_id": event_id, "result": result, "message": message, "scanned_at": scanned_at}


def correct_actual_qty(job_run_id, actual_qty, operator_name):
    try:
        actual_qty = int(actual_qty)
    except (TypeError, ValueError) as exc:
        raise ValueError("完成量必須是非負整數") from exc
    if actual_qty < 0:
        raise ValueError("完成量必須是非負整數")
    db = get_db()
    job = get_job(job_run_id)
    if not job:
        raise ValueError("找不到報工")
    if job["status"] != "finished":
        raise ValueError("只有已完成報工可更正完成量")
    old_qty = job["actual_qty"]
    db.execute("UPDATE job_run SET actual_qty=?,updated_at=? WHERE job_run_id=?", (actual_qty, utc_now(), job_run_id))
    add_audit("correct_actual_qty", "job_run", job_run_id, {"old_actual_qty": old_qty, "new_actual_qty": actual_qty}, operator_name)
    db.commit()


def query_report(filters):
    db = get_db()
    clauses, params = [], []
    status_group = (filters.get("status_group") or "").strip()
    if status_group == "incomplete":
        clauses.append("j.status IN ('created','started')")
    elif status_group == "finished":
        clauses.append("j.status='finished'")
    elif status_group == "cancelled":
        clauses.append("j.status='cancelled'")
    mapping = {"model_code": "j.model_code_snapshot", "machine_no": "COALESCE(j.product_serial,j.machine_no_snapshot)", "status": "j.status", "operator_name": "j.operator_name", "batch_no": "b.batch_no"}
    for key, column in mapping.items():
        value = (filters.get(key) or "").strip()
        if value:
            clauses.append(f"{column} LIKE ?")
            params.append(f"%{value}%")
    for key, op in (("date_from", ">="), ("date_to", "<=")):
        value = (filters.get(key) or "").strip()
        if value:
            clauses.append(f"substr(j.created_at,1,10) {op} ?")
            params.append(value)
    sql = "SELECT j.*, b.batch_no, b.selection_mode, b.created_at AS batch_created_at FROM job_run j JOIN production_batch b ON b.batch_id=j.batch_id"
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    return db.execute(sql + " ORDER BY j.created_at DESC, j.id DESC", params).fetchall()


def export_xlsx(rows, path):
    from copy import copy
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    headers = ["批次號", "job_run識別碼", "機型名稱", "機型代碼", "產品製號", "製作年份", "狀態", "開始時間", "完成時間", "運作時長", "組裝人員", "建立時間", "最後更新時間"]
    wb = Workbook()
    ws = wb.active
    ws.title = "明細"
    ws.append(headers)
    def excel_value(value):
        if not isinstance(value, str):
            return value
        return "'" + value if value[:1] in {"=", "+", "-", "@"} else value

    for row in rows:
        values = [row["batch_no"], row["job_run_id"], row["model_name_snapshot"], row["model_code_snapshot"], row["product_serial"] or row["machine_no_snapshot"], row["manufacturing_year"], status_label(row["status"]), row["started_at"], row["finished_at"], format_duration(row["started_at"], row["finished_at"], row["status"]), row["operator_name"], row["created_at"], row["updated_at"]]
        ws.append([excel_value(v) for v in values])
    if not rows:
        ws.append(["無資料"] + [None] * (len(headers) - 1))
    for cell in ws[1]:
        header_font = copy(cell.font)
        header_font.bold = True
        cell.font = header_font
    for idx, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(header) + 2, 12), 24)

    by_model = wb.create_sheet("依機型彙總")
    by_model.append(["機型代碼", "機型名稱", "產品台數"])
    model_totals = {}
    for row in rows:
        key = (row["model_code_snapshot"], row["model_name_snapshot"])
        item = model_totals.setdefault(key, 0)
        model_totals[key] = item + 1
    for (code, name), item in model_totals.items():
        by_model.append([code, name, item])

    errors = wb.create_sheet("異常掃描")
    errors.append(["事件ID", "條碼類型", "job_run", "結果", "訊息", "掃描時間", "操作者"])
    for event in get_db().execute("SELECT * FROM scan_event WHERE result NOT IN ('success') ORDER BY scanned_at DESC").fetchall():
        errors.append([event["event_id"], event["barcode_type"], event["job_run_id"], event["result"], event["message"], event["scanned_at"], event["operator_name"]])
    wb.save(path)


def backup_database(target_dir):
    db = get_db()
    target_dir = Path(target_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    source = Path(db.execute("PRAGMA database_list").fetchone()[2])
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    target = target_dir / f"intjr_backup_{timestamp}.db"
    db.execute("PRAGMA wal_checkpoint(TRUNCATE)")
    shutil.copy2(source, target)
    check = sqlite3.connect(target)
    try:
        integrity = check.execute("PRAGMA integrity_check").fetchone()[0]
    finally:
        check.close()
    if integrity != "ok" or target.stat().st_size <= 0:
        raise ValueError("備份驗證失敗")
    add_audit("backup", "database", str(target), {"integrity": integrity})
    db.commit()
    return target


def verify_sqlite_backup(path):
    path = Path(path)
    if not path.is_file() or path.stat().st_size <= 0:
        raise ValueError("備份檔不存在或為空檔")
    check = sqlite3.connect(path)
    try:
        integrity = check.execute("PRAGMA integrity_check").fetchone()[0]
    finally:
        check.close()
    if integrity != "ok":
        raise ValueError("備份檔 integrity_check 失敗")
    return path


def restore_database(source_path, target_path, backup_dir, operator_name="admin"):
    source_path = verify_sqlite_backup(source_path)
    safety_backup = backup_database(backup_dir)
    target_path = Path(target_path)
    temp_target = target_path.with_suffix(target_path.suffix + ".restore.tmp")
    shutil.copy2(source_path, temp_target)
    verify_sqlite_backup(temp_target)
    shutil.copy2(temp_target, target_path)
    temp_target.unlink(missing_ok=True)
    add_audit("restore", "database", str(target_path), {"source": str(source_path), "safety_backup": str(safety_backup)}, operator_name)
    get_db().commit()
    return safety_backup

