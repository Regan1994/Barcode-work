import sqlite3
from pathlib import Path

import pytest

from app import create_app as build_app
from app.db import get_db
from app.services import backup_database, create_assembly_worker, create_batch, create_machine_range, create_model, ensure_barcodes, format_datetime, format_duration, get_job, machine_ids_from_numbers, parse_payload, restore_database, scan, set_assembly_worker_active, set_setting, status_label, verify_sqlite_backup


def setup_data(app):
    with app.app_context():
        model_id = create_model("HR-255", "HR255")
        machine_ids = create_machine_range(model_id, "E", 1, 10, 2, "unknown")
        return model_id, machine_ids


def test_machine_range_and_unknown_semantics(app):
    model_id, machine_ids = setup_data(app)
    with app.app_context():
        rows = get_db().execute("SELECT display_no,number_semantics FROM machine_master ORDER BY number_value").fetchall()
        assert len(machine_ids) == 10
        assert [r["display_no"] for r in rows] == [f"E-{i:02d}" for i in range(1, 11)]
        assert {r["number_semantics"] for r in rows} == {"unknown"}


def test_machine_range_validates_semantics(app):
    model_id, machine_ids = setup_data(app)
    with app.app_context():
        with pytest.raises(ValueError, match="編號語意必須是 unknown、asset 或 serial"):
            create_machine_range(model_id, "X", 1, 1, 2, "future")

        assert get_db().execute("SELECT COUNT(*) AS count FROM job_run").fetchone()["count"] == 0


def test_manual_machine_numbers_validate_selection_and_reject_invalid_values(app):
    model_id, machine_ids = setup_data(app)
    with app.app_context():
        assert machine_ids_from_numbers(model_id, " 01, 03, 06 ") == [machine_ids[0], machine_ids[2], machine_ids[5]]
        for invalid in ("01,01", "01,99", "01,,02", "E-01"):
            with pytest.raises(ValueError):
                machine_ids_from_numbers(model_id, invalid)
        assert get_db().execute("SELECT COUNT(*) AS count FROM production_batch").fetchone()["count"] == 0


def test_manual_machine_numbers_reject_other_model_and_inactive_machines(app):
    model_id, machine_ids = setup_data(app)
    with app.app_context():
        other_model = create_model("FAR-210", "FAR210")
        other_machine = create_machine_range(other_model, "F", 1, 1, 2)[0]
        get_db().execute("UPDATE machine_master SET is_active=0 WHERE id=?", (machine_ids[1],))
        get_db().commit()
        for number in ("F-01", "E-02"):
            with pytest.raises(ValueError, match="機台不存在、未啟用或不屬於所選機型"):
                machine_ids_from_numbers(model_id, number)
        assert other_machine


def test_job_create_manual_selection_creates_expected_batch(client, app):
    model_id, _ = setup_data(app)
    with app.app_context():
        create_assembly_worker("operator")
    form = client.get(f"/jobs/new?model_id={model_id}")
    assert form.status_code == 200
    form_text = form.get_data(as_text=True)
    assert "起始流水號" in form_text
    assert "本批台數" in form_text
    assert "HR-255 — 製號字首：E" in form_text
    for removed in ("工單號", "產品批號", "製程", "計畫數量", "機台數字"):
        assert removed not in form_text
    response = client.post("/jobs", data={
        "model_id": str(model_id),
        "manufacturing_year": "2026",
        "starting_number": "10",
        "product_count": "10",
        "operator_name": "operator",
    })
    assert response.status_code == 302
    with app.app_context():
        jobs = get_db().execute("SELECT product_serial,machine_id,machine_no_snapshot,planned_qty,manufacturing_year FROM job_run ORDER BY id").fetchall()
        assert [row["product_serial"] for row in jobs] == [f"E27-{i}" for i in range(10, 20)]
        assert all(row["machine_id"] is None and row["machine_no_snapshot"] is None and row["planned_qty"] is None and row["manufacturing_year"] == 2026 for row in jobs)


def test_five_axis_stage_is_required_and_standard_four_axis_is_disabled(client, app):
    with app.app_context():
        standard_id = create_model("HR-255", "HR255")
        five_axis_id = create_model("FHR-255CL", "FHR255CL", serial_prefix="F")
        create_assembly_worker("operator")

    standard_form = client.get(f"/jobs/new?model_id={standard_id}").get_data(as_text=True)
    assert '<select name="assembly_stage" disabled>' in standard_form
    assert "不適用（標準四軸）" in standard_form

    five_axis_form = client.get(f"/jobs/new?model_id={five_axis_id}").get_data(as_text=True)
    assert '<select name="assembly_stage" required>' in five_axis_form
    for label in ("旋轉軸", "傾斜軸", "尾座", "組合"):
        assert label in five_axis_form

    missing_stage = client.post("/jobs", data={
        "model_id": str(five_axis_id),
        "manufacturing_year": "2026",
        "starting_number": "1",
        "product_count": "1",
        "operator_name": "operator",
    })
    assert missing_stage.status_code == 302
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM production_batch").fetchone()[0] == 0

    created = client.post("/jobs", data={
        "model_id": str(five_axis_id),
        "manufacturing_year": "2026",
        "starting_number": "1",
        "product_count": "1",
        "operator_name": "operator",
        "assembly_stage": "旋",
    })
    assert created.status_code == 302
    with app.app_context():
        batch_id = get_db().execute("SELECT batch_id FROM production_batch").fetchone()["batch_id"]
        batch = get_db().execute("SELECT assembly_stage FROM production_batch").fetchone()
        assert batch["assembly_stage"] == "旋"
        ensure_barcodes(batch_id)
        barcode = get_db().execute("SELECT payload FROM barcode_record WHERE batch_id=? AND barcode_type='START'", (batch_id,)).fetchone()
        assert parse_payload(barcode["payload"])["assembly_stage"] == "旋"

    preview = client.get(f"/batches/{batch_id}/print").get_data(as_text=True)
    assert "組裝階段：旋" in preview


def test_five_axis_same_serial_can_be_created_once_per_stage(app):
    with app.app_context():
        model_id = create_model("FHR-255CL", "FHR255CL", serial_prefix="SS")
        rotate_batch, _ = create_batch(model_id, 2026, 1, 1, "operator", "旋")
        tilt_batch, _ = create_batch(model_id, 2026, 1, 1, "operator", "傾")

        stages = get_db().execute(
            "SELECT product_serial,assembly_stage FROM job_run ORDER BY id"
        ).fetchall()
        assert [(row["product_serial"], row["assembly_stage"]) for row in stages] == [("SS27-01", "旋"), ("SS27-01", "傾")]
        with pytest.raises(ValueError, match="製號已存在"):
            create_batch(model_id, 2026, 1, 1, "operator", "旋")
        assert rotate_batch != tilt_batch


def test_machine_master_is_not_part_of_new_product_flow(client, app):
    with app.app_context():
        model_id = create_model("HR-255", "HR255")

    home = client.get("/")
    assert home.status_code == 200
    assert "機台主檔" not in home.get_data(as_text=True)

    form = client.get(f"/jobs/new?model_id={model_id}")
    assert form.status_code == 200
    assert "不需要先建立機台主檔" in form.get_data(as_text=True)

    legacy_link = client.get("/machines")
    assert legacy_link.status_code == 302
    assert legacy_link.headers["Location"].endswith("/jobs/new")


def test_settings_hide_legacy_machine_semantics(client):
    page = client.get("/settings")
    assert page.status_code == 200
    text = page.get_data(as_text=True)
    assert "舊機台編號語意" not in text
    assert "number_semantics" not in text
    assert "安全還原" in text


def test_protected_modules_require_password_and_unlock(client, app):
    locked = app.test_client()
    response = locked.get("/reports")
    assert response.status_code == 302
    assert response.headers["Location"] == "/unlock?next=/reports"
    assert locked.get("/jobs/new").status_code == 200
    assert locked.get("/scan").status_code == 200

    wrong = locked.post("/unlock", data={"password": "wrong", "next": "/reports"})
    assert wrong.status_code == 200
    assert "密碼錯誤" in wrong.get_data(as_text=True)

    unlocked = locked.post("/unlock", data={"password": "TJR", "next": "/reports"})
    assert unlocked.status_code == 302
    assert unlocked.headers["Location"].endswith("/reports")
    assert locked.get("/reports").status_code == 200
    locked.get("/")
    assert locked.get("/reports").status_code == 302


def test_settings_can_change_access_password(app):
    locked = app.test_client()
    assert locked.post("/unlock", data={"password": "TJR", "next": "/settings"}).status_code == 302
    changed = locked.post(
        "/settings/password",
        data={"current_password": "TJR", "new_password": "NEWPASS", "confirm_password": "NEWPASS"},
    )
    assert changed.status_code == 302

    fresh = app.test_client()
    old_password = fresh.post("/unlock", data={"password": "TJR", "next": "/settings"})
    assert old_password.status_code == 200
    new_password = fresh.post("/unlock", data={"password": "NEWPASS", "next": "/settings"})
    assert new_password.status_code == 302
    assert new_password.headers["Location"].endswith("/settings")


def test_settings_can_temporarily_disable_access_password(app):
    client = app.test_client()
    assert client.post("/unlock", data={"password": "TJR", "next": "/settings"}).status_code == 302
    disabled = client.post("/settings/password-mode", data={})
    assert disabled.status_code == 302

    fresh = app.test_client()
    assert fresh.get("/reports").status_code == 200
    assert fresh.post("/settings/password-mode", data={"password_enabled": "1"}).status_code == 302
    assert fresh.get("/reports").status_code == 302


def test_scan_has_independent_page(client, app):
    response = client.get("/scan")
    assert response.status_code == 200
    text = response.get_data(as_text=True)
    assert "<h1>報工掃描</h1>" in text
    assert 'id="scan-form"' in text
    assert 'name="operator_name"' not in text
    assert "等待掃描" in text

    model_id, _ = setup_data(app)
    with app.app_context():
        batch_id, _ = create_batch(model_id, 2026, 1, 1, "operator")
        ensure_barcodes(batch_id)
        payload = get_db().execute("SELECT payload FROM barcode_record WHERE batch_id=? AND barcode_type='START'", (batch_id,)).fetchone()["payload"]
    scanned = client.post("/scan", data={"payload": payload})
    assert scanned.status_code == 200
    scanned_text = scanned.get_data(as_text=True)
    assert "success" in scanned_text
    assert f"掃描內容：<code>{payload}</code>" in scanned_text


def test_job_create_invalid_manual_selection_does_not_create_partial_batch(client, app):
    model_id, _ = setup_data(app)
    response = client.post("/jobs", data={
        "model_id": str(model_id),
        "manufacturing_year": "2026",
        "starting_number": "1",
        "product_count": "0",
        "operator_name": "operator",
    })
    assert response.status_code == 302
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM production_batch").fetchone()[0] == 0
        assert get_db().execute("SELECT COUNT(*) FROM job_run").fetchone()[0] == 0


def test_product_serial_duplicate_is_rejected_without_partial_batch(app):
    model_id, _ = setup_data(app)
    with app.app_context():
        create_batch(model_id, 2026, 10, 2, "operator")
        with pytest.raises(ValueError, match="製號已存在"):
            create_batch(model_id, 2026, 11, 2, "operator")
        assert get_db().execute("SELECT COUNT(*) FROM production_batch").fetchone()[0] == 1
        assert get_db().execute("SELECT COUNT(*) FROM job_run").fetchone()[0] == 2


def test_model_prefix_mapping_and_year_marker(app):
    with app.app_context():
        examples = (("HR-255", "HR255", "E"), ("RC-255", "RC255", "ER"), ("AR-210", "AR210", "C"), ("RC-210", "RC210", "CR"))
        for name, code, prefix in examples:
            model_id = create_model(name, code)
            batch_id, _ = create_batch(model_id, 2026, 1, 1, "operator")
            assert get_db().execute("SELECT product_serial FROM job_run WHERE batch_id=?", (batch_id,)).fetchone()[0] == f"{prefix}27-01"


def test_report_filter_print_and_export_use_product_serial(client, app, tmp_path):
    from app.services import export_xlsx, query_report
    from openpyxl import load_workbook

    model_id, _ = setup_data(app)
    with app.app_context():
        batch_id, _ = create_batch(model_id, 2026, 10, 2, "operator")
        ensure_barcodes(batch_id)
        rows = query_report({"machine_no": "E27-11"})
        assert len(rows) == 1 and rows[0]["product_serial"] == "E27-11"
        output = tmp_path / "report.xlsx"
        export_xlsx(query_report({}), output)
        workbook = load_workbook(output, read_only=True)
        assert workbook.sheetnames == ["明細"]
        assert [cell.value for cell in workbook["明細"][1]] == ["機型", "製號", "組裝階段", "狀態", "開始時間", "結束時間", "時長", "組裝人員"]
        assert workbook["明細"]["B2"].value in {"E27-10", "E27-11"}

    preview = client.get(f"/batches/{batch_id}/print").get_data(as_text=True)
    assert '<strong>HR-255</strong>' in preview
    assert 'class="thermal-details"' in preview
    assert "E27-10~E27-11" in preview
    assert '<span class="thermal-count">2台</span>' in preview
    assert '<span class="thermal-worker">operator</span>' in preview
    assert preview.index('<strong>HR-255</strong>') < preview.index('<span class="thermal-count">2台</span>') < preview.index('class="thermal-details"')
    assert preview.count('class="barcode-card"') == 2
    assert "開始" in preview and "結束" in preview
    assert preview.index("<h3>開始</h3>") < preview.index("<h3>結束</h3>")
    assert "製作年份" not in preview

    report = client.get("/reports?model_name=HR-255").get_data(as_text=True)
    assert "<td>HR-255</td>" in report
    assert "<td>HR-255（HR255）</td>" not in report
    assert "組裝階段" in report
    assert "機型名稱" in report


def test_report_shows_five_axis_assembly_stage(client, app, tmp_path):
    from app.services import export_xlsx, query_report
    from openpyxl import load_workbook

    with app.app_context():
        model_id = create_model("FHR-255CL", "FHR255CL", serial_prefix="F")
        batch_id, _ = create_batch(model_id, 2026, 1, 1, "operator", "傾")
        rows = query_report({"model_name": "FHR-255CL"})
        assert rows[0]["assembly_stage"] == "傾"
        output = tmp_path / "five-axis-report.xlsx"
        export_xlsx(rows, output)
        workbook = load_workbook(output, read_only=True)
        assert workbook["明細"]["C2"].value == "傾斜軸"

    report = client.get("/reports?model_name=FHR-255CL").get_data(as_text=True)
    assert "傾斜軸" in report
    assert "FHR-255CL" in report


def test_recent_batches_reprint_reuses_original_batch_and_barcodes(client, app):
    model_id, _ = setup_data(app)
    with app.app_context():
        batch_id, _ = create_batch(model_id, 2026, 10, 2, "operator")
        ensure_barcodes(batch_id)
        batch_no = get_db().execute("SELECT batch_no FROM production_batch WHERE batch_id=?", (batch_id,)).fetchone()[0]
        original_records = get_db().execute("SELECT barcode_type,payload FROM barcode_record WHERE batch_id=? ORDER BY barcode_type", (batch_id,)).fetchall()

    recent = client.get("/batches/recent")
    assert recent.status_code == 200
    recent_text = recent.get_data(as_text=True)
    assert "最近批次" in recent_text
    assert "HR-255" in recent_text
    assert "E27-10~E27-11" in recent_text
    assert batch_no not in recent_text
    assert f"/batches/{batch_id}/print" in recent_text

    reprint = client.get(f"/batches/{batch_id}/print")
    assert reprint.status_code == 200
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM production_batch").fetchone()[0] == 1
        assert get_db().execute("SELECT COUNT(*) FROM job_run WHERE batch_id=?", (batch_id,)).fetchone()[0] == 2
        repeated_records = get_db().execute("SELECT barcode_type,payload FROM barcode_record WHERE batch_id=? ORDER BY barcode_type", (batch_id,)).fetchall()
        assert [(row["barcode_type"], row["payload"]) for row in repeated_records] == [(row["barcode_type"], row["payload"]) for row in original_records]


def test_recent_batch_can_be_deleted_after_backup_and_confirmation(client, app):
    model_id, _ = setup_data(app)
    with app.app_context():
        batch_id, _ = create_batch(model_id, 2026, 10, 2, "operator")
        ensure_barcodes(batch_id)
        start = get_db().execute("SELECT payload FROM barcode_record WHERE batch_id=? AND barcode_type='START'", (batch_id,)).fetchone()["payload"]
        assert scan(start)["result"] == "success"

    recent = client.get("/batches/recent")
    assert f"/batches/{batch_id}/delete" in recent.get_data(as_text=True)
    with client.session_transaction() as session:
        token = session["batch_delete_token"]

    assert client.post(f"/batches/{batch_id}/delete", data={"expected_batch_id": batch_id, "confirmation": "刪除"}).status_code == 302
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM production_batch WHERE batch_id=?", (batch_id,)).fetchone()[0] == 1

    response = client.post(f"/batches/{batch_id}/delete", data={
        "csrf_token": token,
        "expected_batch_id": batch_id,
        "confirmation": "刪除",
    })
    assert response.status_code == 302
    with app.app_context():
        db = get_db()
        assert db.execute("SELECT COUNT(*) FROM production_batch WHERE batch_id=?", (batch_id,)).fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM job_run WHERE batch_id=?", (batch_id,)).fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM barcode_record WHERE batch_id=?", (batch_id,)).fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM scan_event").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM audit_log WHERE action='delete' AND entity_type='production_batch' AND entity_id=?", (batch_id,)).fetchone()[0] == 1

    backups = list(Path(app.config["BACKUP_DIRECTORY"]).glob("intjr_backup_*.db"))
    assert len(backups) == 1
    assert verify_sqlite_backup(backups[0]) == backups[0]
    with sqlite3.connect(backups[0]) as original:
        assert original.execute("SELECT COUNT(*) FROM production_batch WHERE batch_id=?", (batch_id,)).fetchone()[0] == 1


def test_long_model_prints_compact_quantity_on_first_line(client, app):
    with app.app_context():
        model_id = create_model("FHR-400C-540-HR400B-FA", "FHR400C540HR400BFA", serial_prefix="F")
        batch_id, _ = create_batch(model_id, 2026, 1, 10, "王小明", "旋")
        ensure_barcodes(batch_id)

    preview = client.get(f"/batches/{batch_id}/print").get_data(as_text=True)
    assert '<div class="thermal-model-line">' in preview
    assert '<strong>FHR-400C-540-HR400B-FA</strong>' in preview
    assert '<span class="thermal-count">10台</span>' in preview
    assert preview.index('<strong>FHR-400C-540-HR400B-FA</strong>') < preview.index('<span class="thermal-count">10台</span>') < preview.index('class="thermal-details"')
    assert '<span class="thermal-worker">王小明</span>' in preview


def test_model_form_uses_name_without_separate_code(client, app):
    form = client.get("/models").get_data(as_text=True)
    assert 'name="model_code"' not in form
    assert "名稱（代碼）" not in form

    response = client.post("/models", data={"model_name": "AR-210", "serial_prefix": "C", "description": "test"})
    assert response.status_code == 200
    with app.app_context():
        model = get_db().execute("SELECT * FROM model_master WHERE model_name='AR-210'").fetchone()
        assert model["model_code"] == "AR-210"
        batch_id, _ = create_batch(model["id"], 2026, 1, 1, "operator")
        assert get_db().execute("SELECT product_serial FROM job_run WHERE batch_id=?", (batch_id,)).fetchone()[0] == "C27-01"

    report = client.get("/reports?model_name=AR-210").get_data(as_text=True)
    assert "<td>AR-210</td>" in report
    assert "C27-01" in report
    client.post("/models", data={"model_name": "AR-210", "serial_prefix": "C"})
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM model_master").fetchone()[0] == 1


def test_renaming_legacy_model_keeps_code_and_existing_report_snapshot(client, app):
    with app.app_context():
        model_id = create_model("HR-255", "HR255")
        create_batch(model_id, 2026, 1, 1, "operator")

    response = client.post(f"/models/{model_id}/update", data={"model_name": "HR-255A", "serial_prefix": "E", "is_active": "1"})
    assert response.status_code == 302
    with app.app_context():
        model = get_db().execute("SELECT model_name,model_code FROM model_master WHERE id=?", (model_id,)).fetchone()
        assert (model["model_name"], model["model_code"]) == ("HR-255A", "HR255")
        old = get_db().execute("SELECT model_name_snapshot,model_code_snapshot FROM job_run").fetchone()
        assert (old["model_name_snapshot"], old["model_code_snapshot"]) == ("HR-255", "HR255")
    assert "<td>HR-255</td>" in client.get("/reports?model_name=HR-255").get_data(as_text=True)
    assert "<td>HR-255</td>" in client.get("/reports?model_code=HR255").get_data(as_text=True)


def test_unused_model_can_be_deleted_after_backup_and_confirmation(client, app):
    with app.app_context():
        model_id = create_model("AR-210", serial_prefix="C")

    page = client.get("/models").get_data(as_text=True)
    assert f'/models/{model_id}/delete' in page
    with client.session_transaction() as session:
        token = session["model_delete_token"]

    for data in (
        {"expected_name": "AR-210"},
        {"csrf_token": token, "expected_name": "other"},
    ):
        assert client.post(f"/models/{model_id}/delete", data=data).status_code == 302
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM model_master").fetchone()[0] == 1
    assert list(Path(app.config["BACKUP_DIRECTORY"]).glob("intjr_backup_*.db")) == []

    response = client.post(f"/models/{model_id}/delete", data={"csrf_token": token, "expected_name": "AR-210"})
    assert response.status_code == 302
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM model_master").fetchone()[0] == 0
        assert get_db().execute("SELECT COUNT(*) FROM audit_log WHERE action='delete' AND entity_type='model_master'").fetchone()[0] == 1
    backups = list(Path(app.config["BACKUP_DIRECTORY"]).glob("intjr_backup_*.db"))
    assert len(backups) == 1
    assert verify_sqlite_backup(backups[0]) == backups[0]
    with sqlite3.connect(backups[0]) as original:
        assert original.execute("SELECT model_name FROM model_master").fetchone()[0] == "AR-210"


def test_model_with_jobs_or_legacy_machines_cannot_be_deleted(client, app):
    with app.app_context():
        used_id = create_model("HR-255", "HR255")
        create_batch(used_id, 2026, 1, 1, "operator")
        legacy_id = create_model("AR-210", "AR210")
        create_machine_range(legacy_id, "C", 1, 1, 2)

    client.get("/models")
    with client.session_transaction() as session:
        token = session["model_delete_token"]
    for model_id, name in ((used_id, "HR-255"), (legacy_id, "AR-210")):
        response = client.post(f"/models/{model_id}/delete", data={"csrf_token": token, "expected_name": name}, follow_redirects=True)
        assert "請改用停用" in response.get_data(as_text=True)
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM model_master").fetchone()[0] == 2
        assert get_db().execute("SELECT COUNT(*) FROM job_run").fetchone()[0] == 1
        assert get_db().execute("SELECT COUNT(*) FROM machine_master").fetchone()[0] == 1
    assert list(Path(app.config["BACKUP_DIRECTORY"]).glob("intjr_backup_*.db")) == []


def test_model_delete_aborts_when_backup_fails(client, app, monkeypatch):
    with app.app_context():
        model_id = create_model("AR-210", serial_prefix="C")
    client.get("/models")
    with client.session_transaction() as session:
        token = session["model_delete_token"]

    def fail_backup(_target_dir):
        raise OSError("備份磁碟不可用")

    monkeypatch.setattr("app.services.backup_database", fail_backup)
    response = client.post(f"/models/{model_id}/delete", data={"csrf_token": token, "expected_name": "AR-210"}, follow_redirects=True)
    assert "沒有刪除機型" in response.get_data(as_text=True)
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM model_master").fetchone()[0] == 1
        assert get_db().execute("SELECT COUNT(*) FROM audit_log WHERE action='delete' AND entity_type='model_master'").fetchone()[0] == 0


def test_v3_migration_preserves_legacy_job_and_foreign_keys(tmp_path):
    path = tmp_path / "legacy.db"
    legacy = sqlite3.connect(path)
    legacy.executescript("""
        CREATE TABLE schema_migrations(version INTEGER PRIMARY KEY, applied_at TEXT NOT NULL);
        INSERT INTO schema_migrations VALUES(1,'old');
        INSERT INTO schema_migrations VALUES(2,'old');
        CREATE TABLE model_master(id INTEGER PRIMARY KEY, model_name TEXT NOT NULL, model_code TEXT NOT NULL UNIQUE, description TEXT, sort_order INTEGER NOT NULL DEFAULT 0, is_active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE TABLE machine_master(id INTEGER PRIMARY KEY, model_id INTEGER NOT NULL REFERENCES model_master(id), display_no TEXT NOT NULL, number_prefix TEXT, number_value INTEGER, number_width INTEGER, number_semantics TEXT NOT NULL DEFAULT 'unknown', is_active INTEGER NOT NULL DEFAULT 1, sort_order INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL, updated_at TEXT NOT NULL, UNIQUE(model_id,display_no));
        CREATE TABLE production_batch(id INTEGER PRIMARY KEY, batch_id TEXT NOT NULL UNIQUE, batch_no TEXT NOT NULL UNIQUE, selection_mode TEXT NOT NULL, work_order_no TEXT, product_lot_no TEXT, process_name TEXT, operator_name TEXT NOT NULL, created_at TEXT NOT NULL, cancelled_at TEXT, cancelled_by TEXT);
        CREATE TABLE job_run(id INTEGER PRIMARY KEY, job_run_id TEXT NOT NULL UNIQUE, batch_id TEXT NOT NULL REFERENCES production_batch(batch_id), model_id INTEGER NOT NULL REFERENCES model_master(id), machine_id INTEGER NOT NULL REFERENCES machine_master(id), model_name_snapshot TEXT NOT NULL, model_code_snapshot TEXT NOT NULL, machine_no_snapshot TEXT NOT NULL, number_semantics_snapshot TEXT NOT NULL DEFAULT 'unknown', work_order_no TEXT, product_lot_no TEXT, process_name TEXT, planned_qty INTEGER NOT NULL CHECK(planned_qty>0), actual_qty INTEGER, status TEXT NOT NULL DEFAULT 'created', started_at TEXT, finished_at TEXT, operator_name TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
        INSERT INTO model_master VALUES(1,'HR-255','HR255',NULL,0,1,'old','old');
        INSERT INTO machine_master VALUES(1,1,'E-01','E',1,2,'unknown',1,0,'old','old');
        INSERT INTO production_batch VALUES(1,'legacy-batch','B-OLD','manual','WO-OLD',NULL,NULL,'old-user','old',NULL,NULL);
        INSERT INTO job_run VALUES(1,'legacy-job','legacy-batch',1,1,'HR-255','HR255','E-01','unknown','WO-OLD',NULL,NULL,5,3,'finished','start-old','finish-old','old-user','old','old');
    """)
    legacy.commit()
    legacy.close()

    app = build_app({"TESTING": True, "DATABASE": str(path), "BACKUP_DIRECTORY": str(tmp_path / "backups"), "SECRET_KEY": "test"})
    with app.app_context():
        old = get_db().execute("SELECT * FROM job_run WHERE job_run_id='legacy-job'").fetchone()
        assert old["machine_no_snapshot"] == "E-01"
        assert old["planned_qty"] == 5 and old["actual_qty"] == 3
        assert old["product_serial"] is None and old["machine_id"] == 1
        assert get_db().execute("SELECT serial_prefix FROM model_master WHERE id=1").fetchone()[0] == "E"
        assert get_db().execute("PRAGMA foreign_key_check").fetchall() == []
        assert next(row["notnull"] for row in get_db().execute("PRAGMA table_info(job_run)") if row["name"] == "machine_id") == 0
        assert next(row["notnull"] for row in get_db().execute("PRAGMA table_info(job_run)") if row["name"] == "machine_no_snapshot") == 0
        assert get_db().execute("SELECT MAX(version) FROM schema_migrations").fetchone()[0] == 7


def test_payload_round_trip_and_idempotent_scan(app):
    model_id, _ = setup_data(app)
    with app.app_context():
        batch_id, job_ids = create_batch(model_id, 2026, 3, 3, "operator")
        ensure_barcodes(batch_id)
        records = get_db().execute("SELECT * FROM barcode_record WHERE batch_id=? ORDER BY barcode_type", (batch_id,)).fetchall()
        assert len(records) == 2
        assert all(len(record["payload"]) == 15 for record in records)
        start = next(record for record in records if record["barcode_type"] == "START")
        finish = next(record for record in records if record["barcode_type"] == "FINISH")
        parsed = parse_payload(start["payload"])
        assert parsed["type"] == "START"
        assert parsed["scope"] == "batch"
        assert scan(start["payload"])["result"] == "success"
        assert scan(start["payload"])["result"] == "duplicate"
        assert scan(finish["payload"])["result"] == "success"
        assert scan(finish["payload"])["result"] == "duplicate"
        assert all(get_job(job_id)["status"] == "finished" and get_job(job_id)["actual_qty"] is None for job_id in job_ids)
        print_preview = app.test_client().get(f"/batches/{batch_id}/print")
        assert print_preview.status_code == 200
        assert print_preview.data.count(b"barcode-card") == 2


def test_finish_before_start_is_rejected(app):
    model_id, _ = setup_data(app)
    with app.app_context():
        batch_id, job_ids = create_batch(model_id, 2026, 1, 1, "operator")
        ensure_barcodes(batch_id)
        finish = get_db().execute("SELECT payload FROM barcode_record WHERE batch_id=? AND barcode_type='FINISH'", (batch_id,)).fetchone()
        assert scan(finish["payload"])["result"] == "invalid"
        assert get_job(job_ids[0])["status"] == "created"


def test_scan_api(client, app):
    model_id, _ = setup_data(app)
    with app.app_context():
        batch_id, job_ids = create_batch(model_id, 2026, 1, 1, "operator")
        ensure_barcodes(batch_id)
        start = get_db().execute("SELECT payload FROM barcode_record WHERE batch_id=? AND barcode_type='START'", (batch_id,)).fetchone()["payload"]
    response = client.post("/api/scan", json={"payload": start})
    assert response.status_code == 200
    assert response.json["result"] == "success"


def test_dashboard_status_cards_and_report_duration(client, app):
    model_id, _ = setup_data(app)
    with app.app_context():
        batch_id, _ = create_batch(model_id, 2026, 1, 1, "operator")
        ensure_barcodes(batch_id)
        start = get_db().execute("SELECT payload FROM barcode_record WHERE batch_id=? AND barcode_type='START'", (batch_id,)).fetchone()["payload"]

        assert status_label("created") == "未完成"
        assert status_label("finished") == "已完成"
        assert format_duration(None, None, "created") == "-"
        assert format_duration("2026-10-16T08:00:00+00:00", "2026-10-17T09:05:06+00:00", "finished") == "1日 01:05:06"
        assert format_datetime("2026-10-19T03:15:20+00:00") == "2026-10-19 11:15:20"

    home = client.get("/")
    assert home.status_code == 200
    assert b"href=\"/reports?status_group=incomplete\"" in home.data
    home_text = home.get_data(as_text=True)
    assert 'href="/scan"' in home_text and "報工掃描" in home_text
    assert "未完成" in home_text
    assert "已完成" in home_text

    with client.session_transaction() as session:
        session["admin_unlocked"] = True
    incomplete = client.get("/reports?status_group=incomplete")
    assert incomplete.status_code == 200
    assert "E27-01" in incomplete.get_data(as_text=True)
    assert "未完成" in incomplete.get_data(as_text=True)

    with app.app_context():
        assert scan(start)["result"] == "success"
    started = client.get("/reports?status_group=incomplete")
    assert "進行中" in started.get_data(as_text=True)


def test_report_status_group_filters_created_and_finished(client, app):
    model_id, _ = setup_data(app)
    with app.app_context():
        batch_id, _ = create_batch(model_id, 2026, 1, 1, "operator")
        ensure_barcodes(batch_id)
        records = get_db().execute("SELECT barcode_type,payload FROM barcode_record WHERE batch_id=?", (batch_id,)).fetchall()
        payloads = {row["barcode_type"]: row["payload"] for row in records}
        assert scan(payloads["START"])["result"] == "success"
        assert scan(payloads["FINISH"])["result"] == "success"

    assert "無資料" in client.get("/reports?status_group=incomplete").get_data(as_text=True)
    finished = client.get("/reports?status_group=finished")
    assert "已完成" in finished.get_data(as_text=True)
    assert "E27-01" in finished.get_data(as_text=True)


def test_report_operator_dropdown_includes_inactive_and_historical_names(client, app):
    with app.app_context():
        model_id = create_model("HR-255", "HR255")
        create_assembly_worker("王小明")
        create_assembly_worker("王小明二")
        inactive_id = create_assembly_worker("李小華")
        for number, name in enumerate(("王小明", "王小明二", "李小華", "歷史員"), start=1):
            create_batch(model_id, 2026, number, 1, name)
        set_assembly_worker_active(inactive_id, False)

    page = client.get("/reports").get_data(as_text=True)
    assert '<select name="operator_name">' in page
    assert '<option value="">全部</option>' in page
    for name in ("王小明", "王小明二", "李小華", "歷史員"):
        assert f'<option value="{name}"' in page

    filtered = client.get("/reports", query_string={"operator_name": "王小明"}).get_data(as_text=True)
    assert '<option value="王小明" selected>王小明</option>' in filtered
    assert "E27-01" in filtered
    assert "E27-02" not in filtered
    assert "共 1 筆產品" in filtered


def test_backup_and_restore_are_verified(app, tmp_path):
    model_id, machine_ids = setup_data(app)
    with app.app_context():
        backup = backup_database(str(tmp_path / "backups"))
        assert verify_sqlite_backup(backup) == backup
        safety = restore_database(backup, app.config["DATABASE"], app.config["BACKUP_DIRECTORY"], "admin")
        assert safety.is_file()


def test_debug_clear_rejects_missing_confirmation_and_master_dependencies(client, app):
    with app.app_context():
        model_id = create_model("HR-255", "HR255")
        create_assembly_worker("王小明")
        create_batch(model_id, 2026, 1, 1, "王小明")

    page = client.get("/debug/clear")
    assert page.status_code == 200
    assert "1 批" in page.get_data(as_text=True)
    with client.session_transaction() as session:
        token = session["debug_clear_token"]

    for data in (
        {"items": "production", "confirmation": "清除"},
        {"items": "production", "confirmation": "wrong", "csrf_token": token},
        {"items": "models", "confirmation": "清除", "csrf_token": token},
    ):
        assert client.post("/debug/clear", data=data).status_code == 302
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM job_run").fetchone()[0] == 1
        assert get_db().execute("SELECT COUNT(*) FROM model_master").fetchone()[0] == 1
        assert get_db().execute("SELECT COUNT(*) FROM audit_log WHERE action='debug_clear'").fetchone()[0] == 0

    with app.app_context():
        scan("invalid")
    assert client.post("/debug/clear", data={"items": "scans", "confirmation": "清除", "csrf_token": token}).status_code == 302
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM scan_event").fetchone()[0] == 0
        assert get_db().execute("SELECT COUNT(*) FROM job_run").fetchone()[0] == 1


def test_debug_clear_selected_data_has_recoverable_backup(client, app):
    with app.app_context():
        model_id = create_model("HR-255", "HR255")
        create_machine_range(model_id, "E", 1, 1, 2)
        create_assembly_worker("王小明")
        batch_id, _ = create_batch(model_id, 2026, 1, 2, "王小明")
        ensure_barcodes(batch_id)
        scan("invalid")

    client.get("/debug/clear")
    with client.session_transaction() as session:
        token = session["debug_clear_token"]
    response = client.post("/debug/clear", data={
        "items": ["production", "workers", "models"],
        "confirmation": "清除",
        "csrf_token": token,
    })
    assert response.status_code == 302
    with app.app_context():
        db = get_db()
        for table in ("production_batch", "job_run", "barcode_record", "scan_event", "assembly_worker", "model_master", "machine_master"):
            assert db.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM app_settings").fetchone()[0] == 4
        assert db.execute("PRAGMA foreign_key_check").fetchall() == []
        log = db.execute("SELECT details FROM audit_log WHERE action='debug_clear'").fetchone()
        assert log is not None

    backups = list(Path(app.config["BACKUP_DIRECTORY"]).glob("intjr_backup_*.db"))
    assert len(backups) == 1
    assert verify_sqlite_backup(backups[0]) == backups[0]
    with sqlite3.connect(backups[0]) as original:
        assert original.execute("SELECT COUNT(*) FROM job_run").fetchone()[0] == 2
        assert original.execute("SELECT COUNT(*) FROM model_master").fetchone()[0] == 1

